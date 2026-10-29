# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Export every work item (story) of a project as a Word document, an Excel sheet or a CSV file — like Jira's
Export menu. Works the same for projects cloned from Jira (their Jira keys and fields come along) and new ones."""

# Python imports
import csv
import io
from collections import defaultdict
from html import escape

# Third party imports
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Django imports
from django.utils import timezone

# Module imports
from plane.db.models import (
    CustomField,
    CustomFieldValue,
    CycleIssue,
    Issue,
    IssueAssignee,
    IssueComment,
    IssueLabel,
    JiraEntityMap,
    ModuleIssue,
    User,
)

EXPORT_FORMATS = ("word", "excel", "csv")
PRIORITY_LABELS = {"urgent": "Urgent", "high": "High", "medium": "Medium", "low": "Low", "none": "None"}


def _person(user):
    if user is None:
        return ""
    full_name = f"{user.first_name or ''} {user.last_name or ''}".strip()
    return full_name or user.display_name or user.email


def _date(value):
    return value.strftime("%Y-%m-%d") if value else ""


def _datetime(value):
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M") if value else ""


def _custom_value(field, value, people):
    if value in (None, "", []):
        return ""
    if field.field_type == "checkbox":
        return "Yes" if value else "No"
    values = value if isinstance(value, list) else [value]
    if field.field_type == "user":
        values = [people.get(str(user_id), str(user_id)) for user_id in values]
    return ", ".join(str(item) for item in values)


def collect_work_items(project, user_id=None):
    """The work items of the project in key order, each as a dict of its exported fields.

    `user_id` limits the export to the work items that user created (clients of a project that hides the others).
    Returns (custom field names, has Jira keys, items).
    """
    issues = Issue.issue_objects.filter(project=project).select_related(
        "state", "type", "parent", "milestone", "created_by", "estimate_point"
    )
    if user_id is not None:
        issues = issues.filter(created_by_id=user_id)
    issues = list(issues.order_by("sequence_id"))
    issue_ids = [issue.id for issue in issues]

    assignees = defaultdict(list)
    for record in IssueAssignee.objects.filter(issue_id__in=issue_ids).select_related("assignee"):
        assignees[record.issue_id].append(_person(record.assignee))
    labels = defaultdict(list)
    for record in IssueLabel.objects.filter(issue_id__in=issue_ids).select_related("label"):
        labels[record.issue_id].append(record.label.name)
    modules = defaultdict(list)
    for record in ModuleIssue.objects.filter(issue_id__in=issue_ids).select_related("module"):
        modules[record.issue_id].append(record.module.name)
    sprints = {
        record.issue_id: record.cycle.name
        for record in CycleIssue.objects.filter(issue_id__in=issue_ids).select_related("cycle")
    }
    comments = defaultdict(list)
    for comment in IssueComment.objects.filter(issue_id__in=issue_ids).select_related("actor").order_by("created_at"):
        comments[comment.issue_id].append(
            {"author": _person(comment.actor), "at": _datetime(comment.created_at), "html": comment.comment_html or ""}
        )
    jira_keys = dict(
        JiraEntityMap.objects.filter(
            workspace_id=project.workspace_id, entity_type="issue", plane_id__in=issue_ids
        ).values_list("plane_id", "jira_key")
    )

    fields = list(CustomField.objects.filter(project=project).order_by("sort_order", "created_at"))
    people = {
        str(user.id): _person(user)
        for user in User.objects.filter(
            id__in={
                user_id
                for value in CustomFieldValue.objects.filter(
                    issue_id__in=issue_ids, field__field_type="user"
                ).values_list("value", flat=True)
                for user_id in (value if isinstance(value, list) else [value])
                if user_id
            }
        )
    }
    custom_values = defaultdict(dict)
    fields_by_id = {field.id: field for field in fields}
    for record in CustomFieldValue.objects.filter(issue_id__in=issue_ids, field_id__in=fields_by_id):
        field = fields_by_id[record.field_id]
        custom_values[record.issue_id][field.name] = _custom_value(field, record.value, people)

    identifier = project.identifier
    items = []
    for issue in issues:
        parent = issue.parent
        items.append(
            {
                "key": f"{identifier}-{issue.sequence_id}",
                "jira_key": jira_keys.get(issue.id, ""),
                "summary": issue.name,
                "type": issue.type.name if issue.type else "",
                "status": issue.state.name if issue.state else "",
                "priority": PRIORITY_LABELS.get(issue.priority, issue.priority or ""),
                "assignees": ", ".join(assignees[issue.id]),
                "reporter": _person(issue.created_by),
                "parent": f"{identifier}-{parent.sequence_id} {parent.name}" if parent else "",
                "milestone": issue.milestone.name if issue.milestone else "",
                "modules": ", ".join(modules[issue.id]),
                "sprint": sprints.get(issue.id, ""),
                "labels": ", ".join(labels[issue.id]),
                "estimate": issue.estimate_point.value if issue.estimate_point else "",
                "start_date": _date(issue.start_date),
                "due_date": _date(issue.target_date),
                "created": _datetime(issue.created_at),
                "updated": _datetime(issue.updated_at),
                "resolved": _datetime(issue.completed_at),
                "custom": custom_values[issue.id],
                "description_html": issue.description_html or "",
                "description_text": issue.description_stripped or "",
                "comments": comments[issue.id],
            }
        )
    return [field.name for field in fields], bool(jira_keys), items


# (label, item key) of the standard columns, in Jira's order
COLUMNS = [
    ("Key", "key"),
    ("Jira key", "jira_key"),
    ("Summary", "summary"),
    ("Type", "type"),
    ("Status", "status"),
    ("Priority", "priority"),
    ("Assignees", "assignees"),
    ("Reporter", "reporter"),
    ("Feature / Parent", "parent"),
    ("Milestone", "milestone"),
    ("Modules", "modules"),
    ("Sprint", "sprint"),
    ("Labels", "labels"),
    ("Estimate", "estimate"),
    ("Start date", "start_date"),
    ("Due date", "due_date"),
    ("Created", "created"),
    ("Updated", "updated"),
    ("Resolved", "resolved"),
]


def _columns(has_jira_keys):
    return [column for column in COLUMNS if has_jira_keys or column[1] != "jira_key"]


def _merge_matching_custom_fields(custom_names, items):
    """A custom field named like a standard column (Jira's "Reporter") fills that column instead of repeating it.

    Its value wins: it is the one from Jira, also for people who have no account here.
    """
    standard = {label.lower(): key for label, key in COLUMNS}
    merged = [name for name in custom_names if name.lower() in standard]
    for item in items:
        for name in merged:
            value = item["custom"].pop(name, "")
            if value:
                item[standard[name.lower()]] = value
    return [name for name in custom_names if name not in merged]


def _rows(custom_names, has_jira_keys, items):
    columns = _columns(has_jira_keys)
    header = [label for label, _ in columns] + custom_names + ["Description", "Comments"]
    rows = [
        [item[key] for _, key in columns]
        + [item["custom"].get(name, "") for name in custom_names]
        + [item["description_text"], len(item["comments"])]
        for item in items
    ]
    return header, rows


def export_csv(custom_names, has_jira_keys, items):
    header, rows = _rows(custom_names, has_jira_keys, items)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    # a byte order mark, so Excel opens the file as UTF-8
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def export_excel(project, custom_names, has_jira_keys, items):
    header, rows = _rows(custom_names, has_jira_keys, items)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Work items"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, label in enumerate(header, start=1):
        wide = label in ("Summary", "Description", "Feature / Parent")
        sheet.column_dimensions[get_column_letter(index)].width = 60 if wide else 18
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


WORD_STYLE = """
body { font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #172B4D; }
h1 { font-size: 20pt; margin: 0 0 4pt; }
h2 { font-size: 14pt; margin: 18pt 0 6pt; border-bottom: 1px solid #DFE1E6; padding-bottom: 4pt; }
h3 { font-size: 11pt; margin: 10pt 0 4pt; }
.meta { color: #5E6C84; font-size: 9pt; margin-bottom: 12pt; }
table.fields { border-collapse: collapse; width: 100%; }
table.fields td { border: 1px solid #DFE1E6; padding: 3pt 6pt; vertical-align: top; font-size: 10pt; }
table.fields td.label { background: #F4F5F7; font-weight: bold; width: 28%; }
.comment { border-left: 3px solid #DFE1E6; padding-left: 8pt; margin: 6pt 0; }
.comment .by { color: #5E6C84; font-size: 9pt; }
"""


def export_word(project, custom_names, has_jira_keys, items):
    """A Word document (HTML that Word opens, the way Jira exports to Word): one section per work item."""
    columns = [column for column in _columns(has_jira_keys) if column[1] not in ("key", "summary")]
    parts = [
        '<html xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w="urn:schemas-microsoft-com:office:word"'
        ' xmlns="http://www.w3.org/TR/REC-html40"><head><meta charset="utf-8">',
        f"<title>{escape(project.name)} work items</title><style>{WORD_STYLE}</style></head><body>",
        f"<h1>{escape(project.name)}</h1>",
        f'<div class="meta">{len(items)} work items · exported {_datetime(timezone.now())}</div>',
    ]
    for item in items:
        parts.append(f"<h2>[{escape(item['key'])}] {escape(item['summary'])}</h2>")
        parts.append('<table class="fields">')
        values = [(label, item[key]) for label, key in columns] + [
            (name, item["custom"].get(name, "")) for name in custom_names
        ]
        for label, value in values:
            if value in ("", None) or (label == "Jira key" and value == item["key"]):
                continue
            parts.append(f'<tr><td class="label">{escape(label)}</td><td>{escape(str(value))}</td></tr>')
        parts.append("</table>")
        if item["description_html"].strip() and item["description_text"].strip():
            # the description is stored sanitized when the work item is saved
            parts.append(f"<h3>Description</h3>{item['description_html']}")
        if item["comments"]:
            parts.append(f"<h3>Comments ({len(item['comments'])})</h3>")
            for comment in item["comments"]:
                parts.append(
                    f'<div class="comment"><div class="by">{escape(comment["author"])} · {escape(comment["at"])}</div>'
                    f"{comment['html']}</div>"
                )
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def export_work_items(project, export_format, user_id=None):
    """Returns (file bytes, content type, file name) of the project's work items in the given format."""
    custom_names, has_jira_keys, items = collect_work_items(project, user_id=user_id)
    custom_names = _merge_matching_custom_fields(custom_names, items)
    stamp = timezone.localtime().strftime("%Y-%m-%d")
    name = f"{project.identifier}-work-items-{stamp}"
    if export_format == "word":
        return export_word(project, custom_names, has_jira_keys, items), "application/msword", f"{name}.doc"
    if export_format == "excel":
        return (
            export_excel(project, custom_names, has_jira_keys, items),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{name}.xlsx",
        )
    return export_csv(custom_names, has_jira_keys, items), "text/csv; charset=utf-8", f"{name}.csv"
