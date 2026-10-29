# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""The user migration sheet of a workspace: every member with their projects and the managers of those projects,
and the columns the managers fill in — migrate (yes / no), migrate to whom and delete (yes / no) — to clean up Jira.
A manager filters the Project Managers column by their name to see their people."""

# Python imports
import io
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor

# Third party imports
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Django imports
from django.utils import timezone

# Module imports
from plane.db.models import JiraEntityMap, ProjectMember, User, WorkspaceMember
from plane.utils.jira.client import JiraError

HYPERLINK_DOMAINS = ("hyperlinkinfosystem.com", "hyperlinkinfosystem.net.in")
HYPERLINK_TYPE = "Hyperlink Infosystem"
ADMIN = 20

# (header, width); the last three are filled in by the managers
COLUMNS = (
    ("Name", 28),
    ("Email Address", 36),
    ("Type", 22),
    ("Status", 14),
    ("Projects", 40),
    ("Project Managers", 30),
    ("Project (Manager)", 50),
    ("Migrate? (Yes/No) *", 18),
    ("Migrate To *", 40),
    ("Delete? (Yes/No) *", 18),
)
MIGRATE_COLUMN, MIGRATE_TO_COLUMN, DELETE_COLUMN = "H", "I", "J"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
REQUIRED_HEADER_FILL = PatternFill("solid", fgColor="C65911")
REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")
MISSING_FILL = PatternFill("solid", fgColor="F8CBAD")


def _person(user):
    full_name = f"{user.first_name or ''} {user.last_name or ''}".strip()
    return full_name or user.display_name or user.email


def user_type(email):
    """Hyperlink Infosystem for the company's own email domains, empty for everyone else (clients, vendors…)."""
    domain = (email or "").rsplit("@", 1)[-1].strip().lower()
    return HYPERLINK_TYPE if domain in HYPERLINK_DOMAINS else ""


def _project_managers(workspace_id):
    """project id -> manager names: the project lead (the Jira project lead for Jira clones), else its admins."""
    memberships = (
        ProjectMember.objects.filter(workspace_id=workspace_id, is_active=True, member__is_bot=False)
        .select_related("project", "project__project_lead", "member")
        .order_by("project__name")
    )
    managers, admins = {}, defaultdict(list)
    for membership in memberships:
        project = membership.project
        if project.project_lead_id:
            managers[project.id] = [_person(project.project_lead)]
        if membership.role == ADMIN:
            admins[project.id].append(_person(membership.member))
    for project_id, names in admins.items():
        managers.setdefault(project_id, sorted(names))
    return managers


def _row(name, email, status, projects):
    """A row of the sheet; projects is [(project name, [its manager names])]."""
    managers = []
    for _, names in projects:
        managers.extend(manager for manager in names if manager not in managers)
    return {
        "name": name,
        "email": email,
        "type": user_type(email),
        "status": status,
        "projects": [project for project, _ in projects],
        "managers": managers,
        "project_managers": [f"{project} ({', '.join(names) or '—'})" for project, names in projects],
        "project_details": projects,
    }


def _sorted(rows):
    return sorted(rows, key=lambda row: (row["name"].lower(), row["email"].lower()))


def user_migration_rows(workspace_id):
    """One row per member of the workspace (deactivated ones too): name, email, type, projects and their managers."""
    managers = _project_managers(workspace_id)
    projects = defaultdict(list)
    for membership in (
        ProjectMember.objects.filter(workspace_id=workspace_id, is_active=True)
        .select_related("project")
        .order_by("project__name")
    ):
        projects[membership.member_id].append(membership.project)

    members = WorkspaceMember.objects.filter(workspace_id=workspace_id, member__is_bot=False).select_related("member")
    return _sorted(
        _row(
            _person(member.member),
            member.member.email,
            "Active" if member.is_active and member.member.is_active else "Deactivated",
            [(project.name, managers.get(project.id, [])) for project in projects.get(member.member_id, [])],
        )
        for member in members
    )


def _is_manager_role(role_name):
    # "Project Manager", "Ass. Project Manager"
    return "project manager" in (role_name or "").lower()


def jira_user_migration_rows(workspace, client, workers=8, progress=None, directory=None, listed_only=False):
    """One row per person of the Jira site (deactivated ones too, apps left out), read live from Jira: their Jira
    projects (the people in the project's roles, groups expanded, its lead, and everyone who is the assignee or
    reporter of a work item — Jira drops deactivated people from the roles) and the managers of those projects (the
    lead and the people in the Project Manager roles). Jira hides most emails: they come from the directory
    (parse_user_export) or else from an earlier user sync. Only reads: thousands of calls for a big site, so run it
    outside a web request. With listed_only, only the people of the directory (an organization's users export) are
    listed, and the ones it suspended count as deactivated."""
    directory = directory or {}
    people = {
        person["accountId"]: person
        for person in client.users()
        if person.get("accountType") == "atlassian"
        and person.get("accountId")
        and (not listed_only or person["accountId"] in directory)
    }
    jira_projects = sorted(client.projects(), key=lambda project: (project.get("name") or "").lower())

    def read_project(project):
        reader, workers_of_project = client.clone(), set()
        try:
            project_roles = reader.project_roles(project["key"])
        except JiraError:
            project_roles = []
        try:
            for issue in reader.issues(project["key"], fields="assignee,reporter"):
                fields = issue.get("fields") or {}
                workers_of_project.update(
                    (fields.get(name) or {}).get("accountId") for name in ("assignee", "reporter")
                )
        except JiraError:
            pass
        return project["key"], project_roles, workers_of_project

    with ThreadPoolExecutor(max_workers=workers) as pool:
        roles, worked = {}, {}
        for done, (key, project_roles, workers_of_project) in enumerate(pool.map(read_project, jira_projects), start=1):
            roles[key], worked[key] = project_roles, workers_of_project
            if progress and done % 25 == 0:
                progress(f"Read {done}/{len(jira_projects)} projects")

    groups = {}

    def group_people(group_id):
        if group_id not in groups:
            try:
                groups[group_id] = [person.get("accountId") for person in client.group_members(group_id)]
            except JiraError:
                groups[group_id] = []
        return groups[group_id]

    def name_of(account_id):
        listed = (directory.get(account_id) or {}).get("name")
        return listed or (people.get(account_id) or {}).get("displayName") or account_id

    member_projects = defaultdict(list)  # account id -> [(project name, managers)]
    for project in jira_projects:
        lead = (project.get("lead") or {}).get("accountId")
        members, managers = [], [lead] if lead else []
        for role in roles.get(project["key"], []):
            if (role.get("name") or "").lower().startswith("atlassian-addons"):
                continue  # the role Jira apps use
            for actor in role.get("actors") or []:
                if actor.get("actorUser"):
                    account_ids = [actor["actorUser"].get("accountId")]
                elif (actor.get("actorGroup") or {}).get("groupId"):
                    account_ids = group_people(actor["actorGroup"]["groupId"])
                else:
                    continue
                members.extend(account_ids)
                if _is_manager_role(role.get("name")):
                    managers.extend(account_ids)
        manager_names = list(dict.fromkeys(name_of(account_id) for account_id in managers if account_id in people))
        for account_id in dict.fromkeys(
            members + ([lead] if lead else []) + sorted(worked.get(project["key"], []) - {None})
        ):
            if account_id in people:
                member_projects[account_id].append((project.get("name") or project["key"], manager_names))

    stored_emails = {}
    for mapping in JiraEntityMap.objects.filter(workspace=workspace, entity_type="user", jira_id__in=list(people)):
        stored_emails[mapping.jira_id] = mapping.jira_key if "@" in (mapping.jira_key or "") else ""
        if not stored_emails[mapping.jira_id] and mapping.plane_id:
            user = User.objects.filter(pk=mapping.plane_id).only("email").first()
            stored_emails[mapping.jira_id] = user.email if user else ""

    return _sorted(
        _row(
            name_of(account_id),
            (
                person.get("emailAddress")
                or (directory.get(account_id) or {}).get("email")
                or stored_emails.get(account_id)
                or ""
            )
            .strip()
            .lower(),
            "Active"
            if person.get("active") and (directory.get(account_id) or {}).get("active", True)
            else "Deactivated",
            member_projects.get(account_id, []),
        )
        for account_id, person in people.items()
    )


LINE_HEIGHT, MAX_ROW_HEIGHT = 15, 409  # points; 409 is the tallest row Excel allows

# (header, width) of the sheet with one row per person and project
PROJECT_COLUMNS = (
    ("Project", 40),
    ("Project Managers", 34),
    ("Name", 28),
    ("Email Address", 36),
    ("Type", 22),
    ("Status", 14),
)


def _wrapped_lines(value, width):
    """About how many lines the text takes in a column of that width (a character is about one width unit)."""
    if not value:
        return 1
    return sum(max(1, -(-len(line) // max(int(width) - 2, 1))) for line in str(value).split("\n"))


def _header(sheet, columns):
    for index, (header, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_HEADER_FILL if header.endswith("*") else HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_user_projects(workbook, rows):
    """One row per person and project, so every project shows on its own line and filters on its own."""
    sheet = workbook.create_sheet("User Projects")
    _header(sheet, PROJECT_COLUMNS)
    row_index = 1
    for row in sorted(rows, key=lambda row: row["name"].lower()):
        for project, managers in sorted(row["project_details"], key=lambda detail: detail[0].lower()) or [("", [])]:
            row_index += 1
            values = (project, ", ".join(managers), row["name"], row["email"], row["type"], row["status"])
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=column, value=value).alignment = Alignment(
                    vertical="top", wrap_text=True
                )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(PROJECT_COLUMNS))}{max(row_index, 2)}"


def _write_instructions(workbook):
    sheet = workbook.create_sheet("Instructions")
    sheet.column_dimensions["A"].width = 110
    lines = (
        "How to fill this sheet",
        "1. In the Users sheet, open the filter of the 'Project Managers' column and pick (or search) your name.",
        "   You now see every employee and client in the projects you manage.",
        "2. For every row fill in the three orange columns — they are compulsory:",
        "   • Migrate? — Yes if the user's data has to be migrated, No otherwise.",
        "   • Migrate To — the person the data moves to (pick from the list). Compulsory when Migrate? is Yes.",
        "   • Delete? — Yes if the user has to be deleted from Jira, No otherwise.",
        "3. Empty compulsory cells are shown in red. Send the sheet back once no red cell is left.",
        "",
        "Type is 'Hyperlink Infosystem' for @hyperlinkinfosystem.com and @hyperlinkinfosystem.net.in emails,"
        " else empty.",
        "Status is Active or Deactivated (in Jira, or in the workspace).",
        "Project Managers: for Jira, the project lead and the people in the Project Manager / Ass. Project Manager",
        " roles of each project of the user; in the workspace, the project lead (its admins when it has no lead).",
        "Projects: for Jira, the projects the user is in (roles or lead) or has work items in (assignee / reporter).",
        "Email Address is empty when Jira hides it and no user sync recorded it.",
        "The User Projects sheet has one row per user and project (filter its Project or Project Managers column).",
    )
    for row, line in enumerate(lines, start=1):
        sheet.cell(row=row, column=1, value=line).font = Font(bold=row == 1, size=14 if row == 1 else 11)


def build_user_migration_sheet(workspace, rows=None):
    """The .xlsx file (bytes, content type, file name) of the user migration sheet of the workspace — of its members,
    or of the given rows (see jira_user_migration_rows)."""
    rows = user_migration_rows(workspace.id) if rows is None else rows
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"

    _header(sheet, COLUMNS)

    for row_index, row in enumerate(rows, start=2):
        values = (
            row["name"],
            row["email"],
            row["type"],
            row["status"],
            "\n".join(row["projects"]),
            "\n".join(row["managers"]),
            "\n".join(row["project_managers"]),
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value).alignment = Alignment(vertical="top", wrap_text=True)
        for column in (MIGRATE_COLUMN, MIGRATE_TO_COLUMN, DELETE_COLUMN):
            sheet[f"{column}{row_index}"].fill = REQUIRED_FILL
        # Excel does not grow the rows of a generated file to their wrapped text: without a height only the first
        # project shows
        lines = max(_wrapped_lines(value, COLUMNS[column][1]) for column, value in enumerate(values))
        sheet.row_dimensions[row_index].height = min(LINE_HEIGHT * lines + 4, MAX_ROW_HEIGHT)

    last_row = max(len(rows) + 1, 2)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    # the active people the data can move to, on a hidden sheet (a list longer than 255 characters needs a range)
    lists = workbook.create_sheet("Lists")
    targets = [row for row in rows if row["status"] == "Active"]
    for index, row in enumerate(targets, start=1):
        lists.cell(row=index, column=1, value=f"{row['name']} <{row['email']}>" if row["email"] else row["name"])
    lists.sheet_state = "hidden"

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False, showErrorMessage=True)
    yes_no.error, yes_no.errorTitle = "Pick Yes or No.", "Compulsory"
    yes_no.add(f"{MIGRATE_COLUMN}2:{MIGRATE_COLUMN}{last_row}")
    yes_no.add(f"{DELETE_COLUMN}2:{DELETE_COLUMN}{last_row}")
    migrate_to = DataValidation(
        type="list", formula1=f"=Lists!$A$1:$A${max(len(targets), 1)}", allow_blank=True, showErrorMessage=True
    )
    migrate_to.error, migrate_to.errorTitle = "Pick the person from the list.", "Migrate To"
    migrate_to.add(f"{MIGRATE_TO_COLUMN}2:{MIGRATE_TO_COLUMN}{last_row}")
    sheet.add_data_validation(yes_no)
    sheet.add_data_validation(migrate_to)

    # compulsory cells left empty turn red (Migrate To only when Migrate? is Yes)
    for column in (MIGRATE_COLUMN, DELETE_COLUMN):
        sheet.conditional_formatting.add(
            f"{column}2:{column}{last_row}", FormulaRule(formula=[f"LEN(TRIM({column}2))=0"], fill=MISSING_FILL)
        )
    sheet.conditional_formatting.add(
        f"{MIGRATE_TO_COLUMN}2:{MIGRATE_TO_COLUMN}{last_row}",
        FormulaRule(formula=[f'AND({MIGRATE_COLUMN}2="Yes",LEN(TRIM({MIGRATE_TO_COLUMN}2))=0)'], fill=MISSING_FILL),
    )

    _write_user_projects(workbook, rows)
    _write_instructions(workbook)

    output = io.BytesIO()
    workbook.save(output)
    file_name = f"{workspace.slug}-user-migration-{timezone.localdate().isoformat()}.xlsx"
    return output.getvalue(), CONTENT_TYPE, file_name
