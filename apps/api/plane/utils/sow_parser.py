# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""
Parser for the SOW / project plan workbook (exported as .xlsx or .pdf).

The workbook has one table per sheet/page, each identified by its first header column:

    Milestone ID   | Milestone | Scope / Acceptance Criteria | Payment % [| Amount]
    Deliverable ID | Deliverable | Description | Milestone ID | Milestone
    Resource ID    | Role | Resource Type | Responsibilities
    Module ID      | Module | Scope
    Feature ID     | Module ID | Module | Feature | Details (from SOW) | Milestone ID
    Task ID        | Feature ID | Feature | Task | Module ID | Task Type | Resource ID | Resource Type | Milestone ID

The parser returns plain dicts and a list of validation errors/warnings; it never touches the database.
"""

# Python imports
import io
import re
from decimal import Decimal, InvalidOperation

# Each table: (key, row id pattern, [(field, header aliases)]) — the field order is the column order
# used when a PDF row has to be split by position.
TABLES = {
    "milestones": {
        "id_pattern": r"M\d+",
        "columns": [
            ("id", ["milestone id"]),
            ("name", ["milestone"]),
            ("scope", ["scope / acceptance criteria", "scope", "acceptance criteria"]),
            ("payment_percentage", ["payment %", "payment"]),
        ],
        "optional_columns": [("amount", ["amount", "payment amount", "milestone amount"])],
    },
    "deliverables": {
        "id_pattern": r"D\d+",
        "columns": [
            ("id", ["deliverable id"]),
            ("name", ["deliverable"]),
            ("description", ["description"]),
            ("milestone_id", ["milestone id"]),
            ("milestone", ["milestone"]),
        ],
        "optional_columns": [],
    },
    "resources": {
        "id_pattern": r"R\d+",
        "columns": [
            ("id", ["resource id"]),
            ("role", ["role"]),
            ("resource_type", ["resource type"]),
            ("responsibilities", ["responsibilities"]),
        ],
        "optional_columns": [],
    },
    "modules": {
        "id_pattern": r"MOD\d+",
        "columns": [
            ("id", ["module id"]),
            ("name", ["module"]),
            ("scope", ["scope"]),
        ],
        "optional_columns": [],
    },
    "features": {
        "id_pattern": r"F\d+",
        "columns": [
            ("id", ["feature id"]),
            ("module_id", ["module id"]),
            ("module", ["module"]),
            ("name", ["feature"]),
            ("details", ["details (from sow)", "details"]),
            ("milestone_id", ["milestone id"]),
        ],
        "optional_columns": [],
    },
    "tasks": {
        "id_pattern": r"T\d+",
        "columns": [
            ("id", ["task id"]),
            ("feature_id", ["feature id"]),
            ("feature", ["feature"]),
            ("name", ["task"]),
            ("module_id", ["module id"]),
            ("task_type", ["task type"]),
            ("resource_id", ["resource id"]),
            ("resource_type", ["resource type"]),
            ("milestone_id", ["milestone id"]),
        ],
        "optional_columns": [("status", ["status"])],
    },
}

# The first header column identifies the table
TABLE_BY_FIRST_HEADER = {
    "milestone id": "milestones",
    "deliverable id": "deliverables",
    "resource id": "resources",
    "module id": "modules",
    "feature id": "features",
    "task id": "tasks",
}

# Page titles of the PDF export
PDF_TABLE_TITLES = [
    ("milestones", "milestones"),
    ("deliverables", "deliverables"),
    ("resources", "resources"),
    ("modules", "modules"),
    ("features", "features"),
    ("task breakdown", "tasks"),
    ("tasks", "tasks"),
]

# Lines that end a table in the PDF export
PDF_TABLE_END = re.compile(r"^(Total\b|Note:|Feature Milestone ID =)")


class SOWParseError(Exception):
    pass


def _normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _join_wrapped(previous, addition):
    """Join a wrapped line of a PDF cell back onto the cell text."""
    if not previous:
        return addition
    if re.search(r"[A-Za-z]-$", previous):
        return previous + addition
    return f"{previous} {addition}"


def _parse_decimal(value, percentage=False):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
        # Excel stores 20% as 0.2
        if percentage and number <= 1:
            number = number * 100
        return number.quantize(Decimal("0.01"))
    text = re.sub(r"[^0-9.\-]", "", str(value))
    if not text:
        return None
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


# ------------------------------------------------------------------ xlsx


def _parse_xlsx(content):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise SOWParseError(f"Could not read the Excel file: {e}")

    tables = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index = None
        table_key = None
        for index, row in enumerate(rows[:25]):
            cells = [_normalize_header(cell) for cell in row]
            first = next((cell for cell in cells if cell), None)
            if first in TABLE_BY_FIRST_HEADER:
                header_index = index
                table_key = TABLE_BY_FIRST_HEADER[first]
                break
        if table_key is None or table_key in tables:
            continue

        schema = TABLES[table_key]
        headers = [_normalize_header(cell) for cell in rows[header_index]]
        column_index = {}
        used = set()
        for field, aliases in schema["columns"] + schema["optional_columns"]:
            for alias in aliases:
                match = next((i for i, h in enumerate(headers) if h == alias and i not in used), None)
                if match is not None:
                    column_index[field] = match
                    used.add(match)
                    break

        missing = [field for field, _ in schema["columns"] if field not in column_index]
        if missing:
            raise SOWParseError(f"Sheet '{sheet.title}' is missing the column(s): {', '.join(missing)}")

        id_regex = re.compile(rf"^{schema['id_pattern']}$")
        records = []
        for row in rows[header_index + 1 :]:
            raw = {field: (row[i] if i < len(row) else None) for field, i in column_index.items()}
            row_id = _clean(raw.get("id"))
            if not id_regex.match(row_id):
                continue
            records.append({field: raw[field] if field in ("payment_percentage", "amount") else _clean(raw[field]) for field in raw})
        tables[table_key] = records
    return tables


# ------------------------------------------------------------------ pdf

_CELL_SPLIT = re.compile(r"\S+(?: \S+)*")


def _line_cells(line):
    """Split a layout line into (start_offset, text) cells separated by 2+ spaces."""
    return [(m.start(), m.group(0)) for m in _CELL_SPLIT.finditer(line)]


def _parse_pdf(content):
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
    except Exception as e:
        raise SOWParseError(f"Could not read the PDF file: {e}")

    tables = {}
    for page in pages:
        lines = page.splitlines()
        table_key = None
        current = None  # (row dict, [column start offsets])
        records = None
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            normalized = _normalize_header(stripped)
            # The page title (or the first header column) identifies which table this page holds
            if table_key is None:
                table_key = next(
                    (key for title, key in PDF_TABLE_TITLES if normalized.startswith(title)),
                    TABLE_BY_FIRST_HEADER.get(" ".join(normalized.split()[:2])),
                )
                if table_key is not None:
                    records = tables.setdefault(table_key, [])
                continue

            if PDF_TABLE_END.match(stripped):
                # Remaining lines on the page are totals / notes
                break

            schema = TABLES[table_key]
            columns = [field for field, _ in schema["columns"]]
            cells = _line_cells(line)
            if not cells:
                continue

            if re.match(rf"^{schema['id_pattern']}$", cells[0][1]) and cells[0][0] < 8:
                if len(cells) != len(columns):
                    raise SOWParseError(
                        f"Could not read row '{cells[0][1]}' in the {table_key} table "
                        f"(found {len(cells)} columns, expected {len(columns)}). "
                        "Please upload the Excel version of the SOW."
                    )
                row = {field: text for field, (_, text) in zip(columns, cells)}
                current = (row, [start for start, _ in cells])
                records.append(row)
                continue

            # Continuation line: each cell belongs to the column whose start offset is nearest
            if current is None:
                # Wrapped header text such as "ID" under "Feature"
                continue
            row, starts = current
            for start, text in cells:
                column = min(range(len(starts)), key=lambda i: abs(starts[i] - start))
                field = columns[column]
                row[field] = _join_wrapped(row[field], text)
    return tables


# ------------------------------------------------------------------ public


def parse_sow(file_name, content):
    """Parse an uploaded SOW file into normalized tables plus validation messages."""
    name = (file_name or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        tables = _parse_xlsx(content)
    elif name.endswith(".pdf"):
        tables = _parse_pdf(content)
    else:
        raise SOWParseError("Unsupported file type. Upload the SOW as .xlsx or .pdf")

    missing_tables = [key for key in TABLES if not tables.get(key)]
    if missing_tables:
        raise SOWParseError(f"The SOW is missing the table(s): {', '.join(missing_tables)}")

    data = {
        "milestones": [
            {
                "id": row["id"],
                "name": row["name"],
                "scope": row["scope"],
                "payment_percentage": _parse_decimal(row.get("payment_percentage"), percentage=True),
                "amount": _parse_decimal(row.get("amount")),
            }
            for row in tables["milestones"]
        ],
        "deliverables": [
            {
                "id": row["id"],
                "name": row["name"],
                "description": row["description"],
                "milestone_id": row["milestone_id"],
            }
            for row in tables["deliverables"]
        ],
        "resources": [
            {
                "id": row["id"],
                "role": row["role"],
                "resource_type": row["resource_type"],
                "responsibilities": row["responsibilities"],
            }
            for row in tables["resources"]
        ],
        "modules": [{"id": row["id"], "name": row["name"], "scope": row["scope"]} for row in tables["modules"]],
        "features": [
            {
                "id": row["id"],
                "module_id": row["module_id"],
                "name": row["name"],
                "details": row["details"],
                "milestone_id": row["milestone_id"],
            }
            for row in tables["features"]
        ],
        "tasks": [
            {
                "id": row["id"],
                "feature_id": row["feature_id"],
                "name": row["name"],
                "module_id": row["module_id"],
                "task_type": row["task_type"],
                "resource_id": row["resource_id"],
                "milestone_id": row["milestone_id"],
                "status": row.get("status", ""),
            }
            for row in tables["tasks"]
        ],
    }
    # Kept only for cross-checking the lookups the workbook repeats on each row
    lookups = {
        "deliverable_milestone_names": {row["id"]: row["milestone"] for row in tables["deliverables"]},
        "feature_module_names": {row["id"]: row["module"] for row in tables["features"]},
        "task_feature_names": {row["id"]: row["feature"] for row in tables["tasks"]},
        "task_resource_types": {row["id"]: row["resource_type"] for row in tables["tasks"]},
    }
    errors, warnings = validate_sow(data, lookups)
    return data, errors, warnings


def validate_sow(data, lookups):
    errors = []
    warnings = []

    def index(key):
        seen = {}
        for row in data[key]:
            if not row["id"]:
                errors.append(f"A row in {key} has no ID")
            elif row["id"] in seen:
                errors.append(f"Duplicate {key[:-1]} ID {row['id']}")
            if not (row.get("name") or row.get("role")):
                errors.append(f"{row['id']} has no name")
            seen[row["id"]] = row
        return seen

    milestones = index("milestones")
    index("deliverables")
    resources = index("resources")
    modules = index("modules")
    features = index("features")
    index("tasks")

    for row in data["deliverables"]:
        milestone = milestones.get(row["milestone_id"])
        if milestone is None:
            errors.append(f"Deliverable {row['id']} refers to unknown milestone {row['milestone_id']}")
        elif _clean(lookups["deliverable_milestone_names"].get(row["id"])) not in ("", milestone["name"]):
            warnings.append(f"Deliverable {row['id']}: milestone name does not match {milestone['id']}")

    for row in data["features"]:
        module = modules.get(row["module_id"])
        if module is None:
            errors.append(f"Feature {row['id']} refers to unknown module {row['module_id']}")
        elif _clean(lookups["feature_module_names"].get(row["id"])) not in ("", module["name"]):
            warnings.append(f"Feature {row['id']}: module name does not match {module['id']}")
        if row["milestone_id"] not in milestones:
            errors.append(f"Feature {row['id']} refers to unknown milestone {row['milestone_id']}")

    for row in data["tasks"]:
        feature = features.get(row["feature_id"])
        if feature is None:
            errors.append(f"Task {row['id']} refers to unknown feature {row['feature_id']}")
        else:
            if _clean(lookups["task_feature_names"].get(row["id"])) not in ("", feature["name"]):
                warnings.append(f"Task {row['id']}: feature name does not match {feature['id']}")
            if row["module_id"] != feature["module_id"]:
                errors.append(
                    f"Task {row['id']} is in module {row['module_id']} but its feature "
                    f"{feature['id']} is in module {feature['module_id']}"
                )
        if row["module_id"] not in modules:
            errors.append(f"Task {row['id']} refers to unknown module {row['module_id']}")
        resource = resources.get(row["resource_id"])
        if resource is None:
            errors.append(f"Task {row['id']} refers to unknown resource {row['resource_id']}")
        elif _clean(lookups["task_resource_types"].get(row["id"])) not in ("", resource["resource_type"]):
            warnings.append(f"Task {row['id']}: resource type does not match {resource['id']}")
        if row["milestone_id"] not in milestones:
            errors.append(f"Task {row['id']} refers to unknown milestone {row['milestone_id']}")
        if not row["task_type"]:
            errors.append(f"Task {row['id']} has no task type")

    total_percentage = sum((row["payment_percentage"] or Decimal(0)) for row in data["milestones"])
    if data["milestones"] and total_percentage != Decimal(100):
        warnings.append(f"Milestone payment percentages add up to {total_percentage}%, not 100%")

    features_with_tasks = {row["feature_id"] for row in data["tasks"]}
    for row in data["features"]:
        if row["id"] not in features_with_tasks:
            warnings.append(f"Feature {row['id']} has no tasks")

    return errors, warnings
