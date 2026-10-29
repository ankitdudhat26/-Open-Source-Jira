# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""The user migration sheet: members, their projects and managers, and the compulsory columns to fill in."""

import io

import pytest
from openpyxl import load_workbook
from rest_framework import status

from plane.db.models import JiraEntityMap, Project, ProjectMember, User, WorkspaceMember
from plane.utils.user_migration_sheet import HYPERLINK_TYPE, jira_user_migration_rows, user_type

URL = "/api/workspaces/{slug}/user-migration-sheet/"


def _member(workspace, email, first_name, role=15):
    user = User.objects.create(email=email, username=email, first_name=first_name, last_name="")
    WorkspaceMember.objects.create(workspace=workspace, member=user, role=role)
    return user


def _rows(response):
    sheet = load_workbook(io.BytesIO(response.content))["Users"]
    header = [cell.value for cell in sheet[1]]
    return header, {row[1]: dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)}


@pytest.mark.unit
def test_user_type_is_hyperlink_only_for_the_company_domains():
    assert user_type("dev@hyperlinkinfosystem.com") == HYPERLINK_TYPE
    assert user_type("Dev@HyperlinkInfosystem.net.in") == HYPERLINK_TYPE
    assert user_type("client@acme.com") == ""
    assert user_type("someone@nothyperlinkinfosystem.com") == ""


class FakeJira:
    """The read only Jira client calls the sheet makes."""

    def __init__(self, users, projects, roles, groups, issues):
        self._users, self._projects, self._roles, self._groups, self._issues = users, projects, roles, groups, issues

    def clone(self):
        return self

    def users(self):
        return self._users

    def projects(self):
        return self._projects

    def project_roles(self, key):
        return self._roles[key]

    def issues(self, key, fields="*all"):
        return self._issues.get(key, [])

    def group_members(self, group_id):
        return [{"accountId": account_id} for account_id in self._groups[group_id]]


def _jira_user(account_id, name, active=True, email="", account_type="atlassian"):
    return {
        "accountId": account_id,
        "displayName": name,
        "active": active,
        "emailAddress": email,
        "accountType": account_type,
    }


def _user_actor(account_id):
    return {"type": "atlassian-user-role-actor", "actorUser": {"accountId": account_id}}


@pytest.mark.unit
@pytest.mark.django_db
def test_jira_rows_have_every_person_with_their_projects_and_managers(workspace):
    client = FakeJira(
        users=[
            _jira_user("pm", "Parth", email="parth@hyperlinkinfosystem.com"),
            _jira_user("dev", "Dev"),
            _jira_user("old", "Old Client", active=False),
            _jira_user("gone", "Gone Dev", active=False),
            _jira_user("bot", "Automation for Jira", account_type="app"),
        ],
        projects=[{"key": "SHOP", "name": "Shop", "lead": {"accountId": "pm"}}, {"key": "BLOG", "name": "Blog"}],
        roles={
            "SHOP": [
                {"name": "Project Manager", "actors": [_user_actor("pm")]},
                {"name": "Developer", "actors": [{"actorGroup": {"groupId": "devs"}}]},
                {"name": "atlassian-addons-project-access", "actors": [_user_actor("bot")]},
            ],
            "BLOG": [{"name": "Client", "actors": [_user_actor("old")]}],
        },
        groups={"devs": ["dev"]},
        # Jira drops deactivated people from the roles: their work items still tie them to the project
        issues={"SHOP": [{"fields": {"assignee": {"accountId": "gone"}, "reporter": None}}]},
    )
    # Jira hides the email of Dev: the one the user sync stored is used
    JiraEntityMap.objects.create(workspace=workspace, entity_type="user", jira_id="dev", jira_key="dev@acme.com")

    directory = {"gone": {"email": "gone@hyperlinkinfosystem.com"}, "pm": {"name": "Parth Mehta", "email": ""}}
    rows = {row["name"]: row for row in jira_user_migration_rows(workspace, client, directory=directory)}

    # the names of the users export win
    assert set(rows) == {"Parth Mehta", "Dev", "Old Client", "Gone Dev"}
    assert rows["Gone Dev"]["projects"] == ["Shop"]
    assert rows["Gone Dev"]["type"] == HYPERLINK_TYPE
    assert rows["Parth Mehta"]["type"] == HYPERLINK_TYPE
    assert rows["Dev"]["email"] == "dev@acme.com"
    assert rows["Dev"]["projects"] == ["Shop"]
    assert rows["Dev"]["managers"] == ["Parth Mehta"]
    assert rows["Old Client"]["status"] == "Deactivated"
    assert rows["Old Client"]["project_managers"] == ["Blog (—)"]


@pytest.mark.contract
class TestUserMigrationSheet:
    @pytest.mark.django_db
    def test_sheet_lists_members_with_projects_and_managers(self, session_client, workspace, create_user):
        manager = _member(workspace, "riya@hyperlinkinfosystem.com", "Riya")
        developer = _member(workspace, "dev@hyperlinkinfosystem.net.in", "Dev")
        client = _member(workspace, "client@acme.com", "Client")
        shop = Project.objects.create(name="Shop", identifier="SHOP", workspace=workspace, project_lead=manager)
        blog = Project.objects.create(name="Blog", identifier="BLOG", workspace=workspace)
        for project, user, role in (
            (shop, manager, 20),
            (shop, developer, 15),
            (shop, client, 5),
            (blog, create_user, 20),
            (blog, developer, 15),
        ):
            ProjectMember.objects.create(project=project, workspace=workspace, member=user, role=role)

        response = session_client.get(URL.format(slug=workspace.slug))

        assert response.status_code == status.HTTP_200_OK
        assert response["Content-Disposition"].endswith('.xlsx"')
        header, rows = _rows(response)
        assert header == [
            "Name",
            "Email Address",
            "Type",
            "Status",
            "Projects",
            "Project Managers",
            "Project (Manager)",
            "Migrate? (Yes/No) *",
            "Migrate To *",
            "Delete? (Yes/No) *",
        ]
        dev = rows["dev@hyperlinkinfosystem.net.in"]
        assert dev["Type"] == HYPERLINK_TYPE
        assert dev["Projects"].split("\n") == ["Blog", "Shop"]
        # Shop has a lead; Blog has none, so its admins manage it
        lead_name = f"{create_user.first_name} {create_user.last_name}".strip()
        assert set(dev["Project Managers"].split("\n")) == {"Riya", lead_name}
        assert "Shop (Riya)" in dev["Project (Manager)"]
        assert rows["client@acme.com"]["Type"] is None
        assert rows["client@acme.com"]["Project Managers"] == "Riya"

    @pytest.mark.django_db
    def test_sheet_has_compulsory_dropdowns_and_a_filter(self, session_client, workspace):
        _member(workspace, "client@acme.com", "Client")

        response = session_client.get(URL.format(slug=workspace.slug))

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook["Users"]
        assert sheet.auto_filter.ref.startswith("A1:J")
        formulas = {str(validation.sqref): validation.formula1 for validation in sheet.data_validations.dataValidation}
        assert '"Yes,No"' in formulas.values()
        assert any(formula.startswith("=Lists!") for formula in formulas.values())
        assert workbook["Lists"].sheet_state == "hidden"
        # Excel does not size generated rows itself
        assert sheet.row_dimensions[2].height
        assert [cell.value for cell in workbook["User Projects"][1]][:3] == ["Project", "Project Managers", "Name"]

    @pytest.mark.django_db
    def test_deactivated_members_are_listed_but_not_offered_to_migrate_to(self, session_client, workspace):
        left = _member(workspace, "left@acme.com", "Left")
        WorkspaceMember.objects.filter(member=left).update(is_active=False)

        response = session_client.get(URL.format(slug=workspace.slug))

        _, rows = _rows(response)
        assert rows["left@acme.com"]["Status"] == "Deactivated"
        targets = [row[0] for row in load_workbook(io.BytesIO(response.content))["Lists"].iter_rows(values_only=True)]
        assert not any("left@acme.com" in target for target in targets)

    @pytest.mark.django_db
    def test_only_workspace_admins_can_download(self, api_client, workspace):
        member = _member(workspace, "member@hyperlinkinfosystem.com", "Member")
        api_client.force_authenticate(user=member)

        response = api_client.get(URL.format(slug=workspace.slug))

        assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.unit
@pytest.mark.django_db
def test_jira_rows_can_be_only_the_people_of_the_users_export(workspace):
    client = FakeJira(
        users=[_jira_user("kept", "Kept"), _jira_user("paused", "Paused"), _jira_user("other", "Other")],
        projects=[],
        roles={},
        groups={},
        issues={},
    )
    directory = {
        "kept": {"email": "kept@hyperlinkinfosystem.com", "active": True},
        "paused": {"email": "paused@acme.com", "active": False},
    }

    rows = {
        row["name"]: row for row in jira_user_migration_rows(workspace, client, directory=directory, listed_only=True)
    }

    assert set(rows) == {"Kept", "Paused"}
    assert rows["Kept"]["email"] == "kept@hyperlinkinfosystem.com"
    assert rows["Paused"]["status"] == "Deactivated"
