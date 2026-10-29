# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Exporting the work items (stories) of a project as Word, Excel or CSV — for new projects and Jira clones."""

import csv
import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from plane.db.models import (
    CustomField,
    CustomFieldValue,
    Issue,
    IssueComment,
    JiraEntityMap,
    Project,
    ProjectMember,
    User,
    WorkspaceMember,
)


@pytest.fixture
def project(db, workspace, create_user):
    project = Project.objects.create(
        name="Export Project", identifier="EXP", workspace=workspace, created_by=create_user
    )
    ProjectMember.objects.create(project=project, member=create_user, workspace=workspace, role=20)
    return project


def _url(workspace, project, export_format):
    return f"/api/workspaces/{workspace.slug}/projects/{project.id}/work-items/export/?file_type={export_format}"


def _csv_rows(response):
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


@pytest.mark.contract
class TestWorkItemExport:
    @pytest.mark.django_db
    def test_word_export_has_every_work_item_with_its_comments(self, session_client, workspace, project, create_user):
        story = Issue.objects.create(name="Login with OTP", project=project, workspace=workspace)
        IssueComment.objects.create(
            issue=story, project=project, workspace=workspace, actor=create_user, comment_html="<p>Ready for QA</p>"
        )
        Issue.objects.create(name="Profile page", project=project, workspace=workspace)

        response = session_client.get(_url(workspace, project, "word"))

        assert response.status_code == status.HTTP_200_OK
        assert response["Content-Type"] == "application/msword"
        assert response["Content-Disposition"].endswith('.doc"')
        html = response.content.decode("utf-8")
        assert f"[EXP-{story.sequence_id}] Login with OTP" in html
        assert "Profile page" in html
        assert "Ready for QA" in html

    @pytest.mark.django_db
    def test_excel_export_lists_the_work_items(self, session_client, workspace, project):
        Issue.objects.create(name="Login with OTP", project=project, workspace=workspace)

        response = session_client.get(_url(workspace, project, "excel"))

        assert response.status_code == status.HTTP_200_OK
        sheet = load_workbook(io.BytesIO(response.content)).active
        header = [cell.value for cell in sheet[1]]
        assert header[:3] == ["Key", "Summary", "Type"]
        assert sheet.cell(row=2, column=2).value == "Login with OTP"

    @pytest.mark.django_db
    def test_jira_clone_exports_jira_keys_and_fields(self, session_client, workspace, project):
        story = Issue.objects.create(name="Checkout", project=project, workspace=workspace)
        JiraEntityMap.objects.create(
            workspace=workspace, entity_type="issue", jira_id="10001", plane_id=story.id, jira_key="SHOP-7"
        )
        team = CustomField.objects.create(project=project, workspace=workspace, name="Team")
        reporter = CustomField.objects.create(project=project, workspace=workspace, name="Reporter")
        CustomFieldValue.objects.create(project=project, workspace=workspace, issue=story, field=team, value="Mobile")
        CustomFieldValue.objects.create(
            project=project, workspace=workspace, issue=story, field=reporter, value="Prem Shah"
        )

        rows = _csv_rows(session_client.get(_url(workspace, project, "csv")))

        assert rows[0]["Jira key"] == "SHOP-7"
        assert rows[0]["Team"] == "Mobile"
        # the Jira reporter fills the Reporter column instead of repeating it
        assert rows[0]["Reporter"] == "Prem Shah"
        assert list(rows[0]).count("Reporter") == 1

    @pytest.mark.django_db
    def test_client_of_a_restricted_project_exports_only_their_own(self, workspace, project, create_user):
        project.guest_view_all_features = False
        project.save()
        client = User.objects.create(email=f"client-{uuid4().hex[:6]}@example.com", username=uuid4().hex[:10])
        WorkspaceMember.objects.create(workspace=workspace, member=client, role=5)
        ProjectMember.objects.create(project=project, member=client, workspace=workspace, role=5)
        team_item = Issue.objects.create(name="Team item", project=project, workspace=workspace)
        client_item = Issue.objects.create(name="Client item", project=project, workspace=workspace)
        # saving sets created_by from the request user, so the authors are set directly
        Issue.objects.filter(pk=team_item.pk).update(created_by=create_user)
        Issue.objects.filter(pk=client_item.pk).update(created_by=client)
        api = APIClient()
        api.force_authenticate(user=client)

        rows = _csv_rows(api.get(_url(workspace, project, "csv")))

        assert [row["Summary"] for row in rows] == ["Client item"]

    @pytest.mark.django_db
    def test_unknown_format_is_rejected(self, session_client, workspace, project):
        response = session_client.get(_url(workspace, project, "pdf"))

        assert response.status_code == status.HTTP_400_BAD_REQUEST
